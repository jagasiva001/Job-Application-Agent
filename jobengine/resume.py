import io
from pathlib import Path

from . import config

ALLOWED_EXT = (".txt", ".md", ".pdf", ".docx")


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", "ignore")
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(data))
        lines = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                lines.extend(c.text for c in row.cells)
        return "\n".join(lines)
    raise ValueError("Unsupported resume type. Use .txt, .pdf or .docx, or paste the text.")


def save_original(data: bytes, filename: str) -> str:
    """Keep the uploaded file so the pre-fill browser can attach it. The stored name is fixed, never user-chosen."""
    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXT:
        raise ValueError("Unsupported resume type.")
    folder = Path(config.DB_PATH).resolve().parent / "uploads"
    folder.mkdir(exist_ok=True)
    for old in folder.glob("resume.*"):
        old.unlink()
    path = folder / f"resume{ext}"
    path.write_bytes(data)
    return str(path)
